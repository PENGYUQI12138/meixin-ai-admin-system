"""M4-B real MariaDB integration tests, exclusively on guarded fictional sites."""
import json
import os
import subprocess
import sys
import unittest
import uuid
from datetime import datetime, timedelta
from decimal import Decimal
from unittest.mock import patch

import frappe
from frappe.utils import get_datetime, now_datetime

from meixin_admin import teacher_payroll as pay
from meixin_admin.teacher_hours import confirm_teaching
from meixin_admin.tests.run import assert_isolated_site


class TestM4B(unittest.TestCase):
    def setUp(self):
        assert_isolated_site()
        frappe.db.rollback()
        frappe.set_user("Administrator")
        self.batch = "TEST-M4B-" + uuid.uuid4().hex[:12]
        self.workers, self.users = [], []
        self.committed = False
        self.next_slot = 0
        self.day = datetime(2020, 1, 15, 8)
        settings = frappe.get_single("MX Settings")
        for field in ("present_rule", "leave_rule", "absent_rule", "other_rule", "session_cancel_rule"):
            settings.set(field, "不课消")
        settings.save()
        self.students = [self.master("MX Student", student_name=f"工资虚构学生{i}",
                                     guardian_phone="00000000000") for i in range(5)]
        self.teacher = self.master("MX Teacher", teacher_name="工资虚构教师")
        self.course = self.master("MX Course", course_name="工资虚构课程", default_duration_minutes=90)
        self.room = self.master("MX Room", room_name="工资虚构教室", capacity=8)
        self.rule = self.new_rule("2020-01-01")

    def tearDown(self):
        for process in self.workers:
            if process.poll() is None:
                process.kill()
            process.wait(timeout=10)
        frappe.set_user("Administrator")
        frappe.db.rollback()
        if not self.committed:
            return
        parent_children = (
            ("MX Teacher Pay Sheet", ("MX Teacher Pay Line", "MX Teacher Pay Coverage")),
            ("MX Teacher Pay Rule", ("MX Teacher Pay Rate",)),
            ("MX Session Execution", ("MX Session Attendance",)),
            ("MX Session", ("MX Session Student",)),
        )
        for parent, children in parent_children:
            names = frappe.get_all(parent, filters={"demo_batch": self.batch}, pluck="name")
            if names:
                for child in children:
                    frappe.db.delete(child, {"parent": ["in", names]})
        for doctype in (
            "MX Teacher Pay Payment", "MX Teacher Pay Adjustment", "MX Teacher Pay Sheet",
            "MX Teacher Pay Rule", "MX Teacher Hour Entry", "MX Lesson Consumption Entry",
            "MX Session Execution", "MX Session", "MX Student", "MX Teacher", "MX Course", "MX Room",
        ):
            names = frappe.get_all(doctype, filters={"demo_batch": self.batch}, pluck="name")
            if names:
                frappe.db.delete("Version", {"ref_doctype": doctype, "docname": ["in", names]})
            frappe.db.delete(doctype, {"demo_batch": self.batch})
        for user in self.users:
            frappe.db.delete("Has Role", {"parent": user})
            frappe.db.delete("User Email", {"parent": user})
            frappe.db.delete("User", {"name": user})
        for field in ("present_rule", "leave_rule", "absent_rule", "other_rule", "session_cancel_rule"):
            frappe.db.set_single_value("MX Settings", field, "未配置")
        frappe.db.commit()

    def master(self, doctype, **fields):
        return frappe.get_doc({"doctype": doctype, "enabled": 1,
                               "demo_batch": self.batch, **fields}).insert()

    def user(self, *roles):
        name = frappe.get_doc({
            "doctype": "User", "email": f"m4b-{uuid.uuid4().hex[:16]}@example.invalid",
            "first_name": "工资隔离用户", "enabled": 1, "send_welcome_email": 0,
            "user_type": "System User", "roles": [{"role": role} for role in roles],
        }).insert().name
        self.users.append(name)
        return name

    def new_rule(self, date, low=100, high=120):
        return frappe.get_doc({
            "doctype": "MX Teacher Pay Rule", "effective_from": date, "currency": "CNY",
            "rates": [{"min_students": 1, "max_students": 4, "amount": low},
                      {"min_students": 5, "max_students": 0, "amount": high}],
            "demo_batch": self.batch,
        }).insert().submit()

    def lesson(self, size=4, present=3, *, start=None, duration=90, seconds=0,
               exception_reason=None):
        if start is None:
            start = self.day + timedelta(hours=2 * self.next_slot)
            self.next_slot += 1
        end = start + timedelta(minutes=duration, seconds=seconds)
        session = frappe.get_doc({
            "doctype": "MX Session", "teacher": self.teacher.name,
            "course": self.course.name, "room": self.room.name,
            "start_at": start, "end_at": end,
            "students": [{"student": student.name} for student in self.students[:size]],
            "demo_batch": self.batch,
        }).insert().submit()
        execution = frappe.get_doc({
            "doctype": "MX Session Execution", "session": session.name,
            "attendance": [{"student": student.name,
                            "attendance_status": "到课" if i < present else "请假"}
                           for i, student in enumerate(self.students[:size])],
        }).insert().submit()
        hour = confirm_teaching(execution.name, start, end,
                                exception_reason=exception_reason or ("无人到课仍授课，须另核实" if not present else None))
        return session, execution, hour

    def line(self, hour, size=4, **fields):
        return {"kind": "课程", "hour_entry": hour, "class_size": size,
                "headcount_basis": "已核对冻结名单包含请假及缺勤学生，教师归属无争议", **fields}

    def fixed(self, *, hours=None, date=None, amount=300, scope="指定", work_key=None):
        return {"kind": "固定", "scope": scope, "hours": [] if scope == "全天" else hours or [],
                "pay_date": date or str(self.day.date()), "amount": amount,
                "exception_reason": "已核对教师与覆盖约定，固定费用",
                "description": "虚构指定教学及备课工作", "work_key": work_key}

    def sheet(self, lines, *, request_id=None):
        return pay.create_sheet(self.teacher.name, "2020-01-01", "2020-12-31",
                                lines, request_id or uuid.uuid4().hex)

    def approved(self, lines):
        sheet = self.sheet(lines)
        self.assertEqual(pay.check_sheet(sheet), sheet)
        self.assertEqual(pay.approve_sheet(sheet), sheet)
        return sheet

    def payment(self, sheet, request_id=None):
        return pay.record_pay_payment(sheet, str(now_datetime()), "现金", "隔离测试实际付款",
                                      request_id or uuid.uuid4().hex)

    def money(self, sheet, field="net_payable"):
        return Decimal(str(pay.get_pay_balance(sheet)[field]))

    def adjustments(self, sheet):
        return frappe.get_all("MX Teacher Pay Adjustment", filters={"sheet": sheet},
                              fields=["name", "docstatus"], order_by="creation, name")

    def commit(self):
        frappe.db.commit()
        self.committed = True

    def worker(self, operation, name, hold, args=None, user=None):
        process = subprocess.Popen(
            [sys.executable, "-m", "meixin_admin.tests.concurrent_worker"],
            stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            text=True, bufsize=1,
        )
        self.workers.append(process)
        process.stdin.write(json.dumps({
            "site": frappe.local.site, "sites_path": os.path.abspath(frappe.local.sites_path),
            "operation": operation, "name": name, "hold": hold, "field": args,
            "user": user,
        }) + "\n")
        process.stdin.flush()
        ready = process.stdout.readline().strip()
        self.assertEqual(ready, "LOCKED" if hold else "READY", process.stderr.read() if not ready else ready)
        return process

    def race(self, first_op, first_name, first_args, second_op=None, second_name=None, second_args=None):
        self.commit()
        first = self.worker(first_op, first_name, True, first_args)
        second = self.worker(second_op or first_op, second_name or first_name, False,
                             second_args if second_args is not None else first_args)
        for process in (second, first):
            process.stdin.write("GO\n")
            process.stdin.flush()
            self.assertEqual(process.stdout.readline().strip(), "ATTEMPT")
        results = []
        for process in (first, second):
            line = process.stdout.readline()
            self.assertTrue(line.startswith("RESULT "), line)
            results.append(json.loads(line.removeprefix("RESULT ")))
            process.wait(timeout=30)
        frappe.db.rollback()
        self.assertTrue(results[0]["ok"], results)
        if not results[1]["ok"]:
            self.assertEqual(results[1]["exception"], "ValidationError", results)
        return results

    def test_01_class_four_attendance_three_and_class_five_attendance_four_total_220(self):
        first = self.lesson(4, 3)[2]
        second = self.lesson(5, 4)[2]
        sheet = self.approved([self.line(first, 4), self.line(second, 5)])
        doc = frappe.get_doc("MX Teacher Pay Sheet", sheet)
        self.assertEqual([Decimal(str(row.amount)) for row in doc.lines], [Decimal(100), Decimal(120)])
        self.assertEqual(self.money(sheet), 220)
        self.assertEqual(pay.get_pay_balance(sheet)["status"], "待付款")

    def test_02_missing_headcount_basis_stays_unchecked(self):
        hour = self.lesson()[2]
        sheet = self.sheet([self.line(hour, headcount_basis="")])
        self.assertEqual(pay.get_pay_balance(sheet)["status"], "待核对")
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(sheet)
        with self.assertRaises(frappe.ValidationError):
            pay.approve_sheet(sheet)

    def test_03_incomplete_roster_can_have_verified_historical_class_size(self):
        hour = self.lesson(4, 3)[2]
        sheet = self.approved([self.line(hour, 5, headcount_basis="历史纸质班册5人，1人请假未列入排课，教师无变更")])
        self.assertEqual(self.money(sheet), 120)

    def test_04_non_ninety_and_ninety_minutes_fifty_nine_seconds_rejected(self):
        for duration, seconds in ((89, 0), (90, 59)):
            with self.subTest(duration=duration, seconds=seconds):
                hour = self.lesson(duration=duration, seconds=seconds)[2]
                sheet = self.sheet([self.line(hour)])
                with self.assertRaises(frappe.ValidationError):
                    pay.check_sheet(sheet)
                self.assertEqual(pay.get_pay_balance(sheet)["status"], "待核对")

    def test_05_no_attendance_requires_special_finance_approved_amount(self):
        hour = self.lesson(present=0)[2]
        sheet = self.sheet([self.line(hour)])
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(sheet)
        pay.update_sheet(sheet, [self.line(hour, kind="特殊", amount=80,
                                          exception_reason="学生全请假，已按教学约定确认备课补偿")])
        pay.check_sheet(sheet)
        pay.approve_sheet(sheet)
        self.assertEqual(self.money(sheet), 80)

    def test_06_cancelled_source_cannot_be_paid_even_as_special(self):
        _, execution, hour = self.lesson()
        execution.cancel()
        for kind in ("课程", "特殊"):
            with self.subTest(kind=kind), self.assertRaises(frappe.ValidationError):
                sheet = self.sheet([self.line(hour, kind=kind, amount=80, exception_reason="不应允许取消课")])
                pay.check_sheet(sheet)
        self.assertEqual(frappe.db.count("MX Teacher Pay Sheet", {"demo_batch": self.batch, "docstatus": 1}), 0)

    def test_07_new_rule_and_new_roster_do_not_rewrite_approved_snapshot(self):
        session, _, hour = self.lesson()
        sheet = self.approved([self.line(hour)])
        before = frappe.get_doc("MX Teacher Pay Sheet", sheet).as_json()
        self.new_rule("2020-01-02", 150, 180)
        # Model an external membership edit: approved payroll must still use its frozen evidence.
        frappe.db.delete("MX Session Student", {"name": session.students[-1].name})
        self.assertEqual(self.money(sheet), 100)
        self.assertEqual(frappe.get_doc("MX Teacher Pay Sheet", sheet).as_json(), before)

    def test_08_rule_change_after_check_requires_recheck_before_approval(self):
        hour = self.lesson()[2]
        sheet = self.sheet([self.line(hour)])
        pay.check_sheet(sheet)
        self.new_rule("2020-01-02", 150, 180)
        with self.assertRaises(frappe.ValidationError):
            pay.approve_sheet(sheet)
        pay.withdraw_sheet(sheet)
        pay.check_sheet(sheet)
        pay.approve_sheet(sheet)
        self.assertEqual(self.money(sheet), 150)

    def test_09_checked_headcount_change_requires_recheck(self):
        hour = self.lesson()[2]
        sheet = self.sheet([self.line(hour)])
        pay.check_sheet(sheet)
        pay.update_sheet(sheet, [self.line(hour, 5)])
        with self.assertRaises(frappe.ValidationError):
            pay.approve_sheet(sheet)
        pay.check_sheet(sheet)
        pay.approve_sheet(sheet)
        self.assertEqual(self.money(sheet), 120)

    def test_10_direct_document_mutation_is_blocked(self):
        hour = self.lesson()[2]
        sheet = self.approved([self.line(hour)])
        doc = frappe.get_doc("MX Teacher Pay Sheet", sheet)
        doc.lines[0].amount = 999
        with self.assertRaises((frappe.PermissionError, frappe.ValidationError)):
            doc.save()
        with self.assertRaises((frappe.PermissionError, frappe.ValidationError)):
            frappe.get_doc("MX Teacher Pay Sheet", sheet).cancel()
        with self.assertRaises((frappe.PermissionError, frappe.ValidationError)):
            frappe.get_doc({"doctype": "MX Teacher Pay Payment", "sheet": sheet,
                            "amount": 100, "paid_at": now_datetime(), "payment_method": "现金"}).insert()
        with self.assertRaises((frappe.PermissionError, frappe.ValidationError)):
            self.rule.cancel()
        self.assertEqual(self.money(sheet), 100)

    def test_11_scoped_fixed_fee_allows_uncovered_course(self):
        first = self.lesson()[2]
        second = self.lesson(5, 4)[2]
        sheet = self.approved([self.fixed(hours=[first], amount=250), self.line(second, 5)])
        self.assertEqual(self.money(sheet), 370)
        duplicate = self.sheet([self.line(first)])
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(duplicate)

    def test_12_whole_day_fixed_fee_blocks_same_day_and_cross_sheet(self):
        first = self.lesson()[2]
        second = self.lesson()[2]
        sheet = self.approved([self.fixed(hours=[first, second], scope="全天", amount=300)])
        self.assertEqual(self.money(sheet), 300)
        duplicate = self.sheet([self.line(second)])
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(duplicate)
        another_day = self.lesson(start=self.day + timedelta(days=1))[2]
        separate = self.approved([self.line(another_day)])
        self.assertEqual(self.money(separate), 100)

    def test_13_later_same_day_course_creates_review_and_cannot_bypass_whole_day(self):
        first = self.lesson()[2]
        sheet = self.approved([self.fixed(hours=[first], scope="全天")])
        new_hour = self.lesson()[2]
        self.assertEqual(len(self.adjustments(sheet)), 1)
        self.assertEqual(pay.get_pay_balance(sheet)["status"], "待处理调整")
        with self.assertRaises(frappe.ValidationError):
            self.payment(sheet)
        duplicate = self.sheet([self.line(new_hour)])
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(duplicate)
        adjustment = self.adjustments(sheet)[0].name
        pay.approve_adjustment(adjustment, 0, "全日固定费已涵盖后补课，不增加金额")
        self.assertEqual(self.money(sheet), 300)

    def test_14_whole_day_changed_before_approval_requires_recheck(self):
        hour = self.lesson()[2]
        sheet = self.sheet([self.fixed(hours=[hour], scope="全天")])
        pay.check_sheet(sheet)
        self.lesson()
        with self.assertRaises(frappe.ValidationError):
            pay.approve_sheet(sheet)

    def test_15_specific_work_key_prevents_cross_sheet_duplicate(self):
        sheet = self.approved([self.fixed(work_key="虚构教案工作-" + self.batch, amount=66)])
        self.assertEqual(self.money(sheet), 66)
        duplicate = self.sheet([self.fixed(work_key="虚构教案工作-" + self.batch, amount=66)])
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(duplicate)

    def test_16_fixed_scope_and_special_reason_are_required(self):
        hour = self.lesson()[2]
        for line in (self.fixed(), self.fixed(hours=[hour], scope=""),
                     self.line(hour, kind="特殊", amount=80, exception_reason="")):
            with self.subTest(line=line), self.assertRaises(frappe.ValidationError):
                sheet = self.sheet([line])
                pay.check_sheet(sheet)

    def test_17_drafts_do_not_claim_and_withdraw_releases_unapproved_source(self):
        hour = self.lesson()[2]
        first = self.sheet([self.line(hour)])
        second = self.sheet([self.line(hour)])
        pay.check_sheet(first)
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(second)
        pay.withdraw_sheet(first)
        pay.check_sheet(second)
        pay.approve_sheet(second)
        self.assertEqual(self.money(second), 100)

    def test_18_cancel_before_approval_blocks_original(self):
        _, execution, hour = self.lesson()
        sheet = self.sheet([self.line(hour)])
        pay.check_sheet(sheet)
        execution.cancel()
        with self.assertRaises(frappe.ValidationError):
            pay.approve_sheet(sheet)
        self.assertEqual(frappe.db.get_value("MX Teacher Pay Sheet", sheet, "docstatus"), 0)
        self.assertEqual(len(self.adjustments(sheet)), 0)

    def test_19_cancel_after_approval_unpaid_blocks_until_net_correction(self):
        _, first_exec, first = self.lesson()
        second = self.lesson(5, 4)[2]
        sheet = self.approved([self.line(first), self.line(second, 5)])
        first_exec.cancel()
        self.assertEqual(len(self.adjustments(sheet)), 1)
        with self.assertRaises(frappe.ValidationError):
            self.payment(sheet)
        adjustment = self.adjustments(sheet)[0].name
        self.assertEqual(pay.approve_adjustment(adjustment, -100, "取消普通课程，冲回该节应付"), adjustment)
        self.assertEqual(self.money(sheet, "total_amount"), 220)
        self.assertEqual(self.money(sheet, "adjustment_amount"), -100)
        self.assertEqual(self.money(sheet), 120)
        payment = self.payment(sheet)
        self.assertEqual(Decimal(str(frappe.db.get_value("MX Teacher Pay Payment", payment, "amount"))), 120)
        self.assertEqual(pay.get_pay_balance(sheet)["status"], "已付款")

    def test_20_paid_source_cancellation_preserves_payment_and_tracks_outstanding(self):
        _, execution, hour = self.lesson()
        sheet = self.approved([self.line(hour)])
        payment = self.payment(sheet)
        original_payment = frappe.get_doc("MX Teacher Pay Payment", payment).as_json()
        execution.cancel()
        adjustment = self.adjustments(sheet)[0].name
        self.assertEqual(pay.get_pay_balance(sheet)["status"], "待处理调整")
        pay.approve_adjustment(adjustment, -100, "已付款课撤销，待财务线下处理")
        self.assertEqual(self.money(sheet, "net_paid"), 100)
        self.assertEqual(self.money(sheet, "outstanding"), -100)
        self.assertEqual(pay.get_pay_balance(sheet)["status"], "待财务处理")
        self.assertEqual(frappe.get_doc("MX Teacher Pay Payment", payment).as_json(), original_payment)
        with self.assertRaises(frappe.ValidationError):
            self.payment(sheet)

    def test_21_fixed_fee_partial_cancellation_never_mechanically_deducts_whole_fee(self):
        _, execution, first = self.lesson()
        second = self.lesson()[2]
        sheet = self.approved([self.fixed(hours=[first, second], amount=300)])
        execution.cancel()
        self.assertEqual(self.money(sheet, "total_amount"), 300)
        self.assertEqual(self.money(sheet, "adjustment_amount"), 0)
        adjustment = self.adjustments(sheet)[0].name
        pay.approve_adjustment(adjustment, -50, "按固定费用约定，仅取消部分授课扣50元")
        self.assertEqual(self.money(sheet), 250)

    def test_22_reconfirmed_execution_revision_cannot_pay_original_lesson_again(self):
        session, execution, hour = self.lesson()
        sheet = self.approved([self.line(hour)])
        execution.cancel()
        amended = frappe.copy_doc(execution)
        amended.amended_from, amended.docstatus = execution.name, 0
        for row in amended.attendance:
            row.consumption_rule_result = None
        amended.insert().submit()
        revised = confirm_teaching(amended.name, session.start_at, session.end_at)
        duplicate = self.sheet([self.line(revised)])
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(duplicate)
        self.assertGreaterEqual(len(self.adjustments(sheet)), 1)

    def test_23_zero_and_negative_net_no_false_payment_or_refund(self):
        for adjustment_amount, expected_status in ((-100, "无需付款"), (-150, "待财务处理")):
            with self.subTest(adjustment_amount=adjustment_amount):
                _, execution, hour = self.lesson()
                sheet = self.approved([self.line(hour)])
                execution.cancel()
                adjustment = self.adjustments(sheet)[0].name
                pay.approve_adjustment(adjustment, adjustment_amount, "财务按约定核定差额")
                self.assertEqual(self.money(sheet), 100 + adjustment_amount)
                self.assertEqual(pay.get_pay_balance(sheet)["status"], expected_status)
                with self.assertRaises(frappe.ValidationError):
                    self.payment(sheet)
                self.assertEqual(frappe.db.count("MX Teacher Pay Payment", {"sheet": sheet}), 0)

    def test_24_payment_idempotency_reversal_and_fresh_registration(self):
        sheet = self.approved([self.line(self.lesson()[2])])
        key, paid_at = uuid.uuid4().hex, str(now_datetime())
        first = pay.record_pay_payment(sheet, paid_at, "现金", "唯一凭据", key)
        self.assertEqual(pay.record_pay_payment(sheet, paid_at, "现金", "唯一凭据", key), first)
        with self.assertRaises(frappe.ValidationError):
            pay.record_pay_payment(sheet, paid_at, "现金", "不同凭据", key)
        with self.assertRaises(frappe.ValidationError):
            self.payment(sheet)
        reverse_key = uuid.uuid4().hex
        reverse = pay.reverse_pay_payment(first, "错误登记，线下尚未实际支付", reverse_key)
        self.assertEqual(pay.reverse_pay_payment(first, "错误登记，线下尚未实际支付", reverse_key), reverse)
        with self.assertRaises(frappe.ValidationError):
            pay.reverse_pay_payment(first, "重复冲正", uuid.uuid4().hex)
        self.assertEqual(self.money(sheet, "net_paid"), 0)
        second = self.payment(sheet)
        self.assertNotEqual(first, second)
        self.assertEqual(frappe.db.count("MX Teacher Pay Payment", {"sheet": sheet}), 3)
        self.assertEqual(self.money(sheet, "net_paid"), 100)
        self.assertEqual(pay.get_pay_balance(sheet)["status"], "已付款")

    def test_25_admin_direct_finance_api_and_rule_write_are_denied(self):
        _, execution, hour = self.lesson()
        sheet = self.sheet([self.line(hour)])
        scheduler = self.user("Meixin Scheduler")
        manager = self.user("Meixin Manager")
        for user in (scheduler, manager):
            frappe.set_user(user)
            pay.check_sheet(sheet)
            with self.assertRaises(frappe.PermissionError):
                pay.approve_sheet(sheet)
            with self.assertRaises(frappe.PermissionError):
                self.payment(sheet)
            with self.assertRaises(frappe.PermissionError):
                self.new_rule("2020-02-01")
        frappe.set_user("Administrator")
        pay.approve_sheet(sheet)
        payment = self.payment(sheet)
        execution.cancel()
        adjustment = self.adjustments(sheet)[0].name
        frappe.set_user(manager)
        with self.assertRaises(frappe.PermissionError):
            pay.approve_adjustment(adjustment, -100, "行政无权批准")
        with self.assertRaises(frappe.PermissionError):
            pay.reverse_pay_payment(payment, "行政无权冲正", uuid.uuid4().hex)

    def test_26_finance_gets_only_source_summary_not_raw_student_data(self):
        session, execution, hour = self.lesson()
        finance = self.user("Meixin Finance")
        frappe.set_user(finance)
        summary = pay.get_source_summary(hour)
        payload = json.dumps(summary, ensure_ascii=False, default=str)
        self.assertIn(self.teacher.teacher_name, payload)
        self.assertIn(self.course.course_name, payload)
        for secret in ("guardian_phone", "attendance", "headcount_basis", "exception_reason",
                       "00000000000", *(student.student_name for student in self.students),
                       *(student.name for student in self.students)):
            self.assertNotIn(secret, payload)
        for doctype, name in (("MX Student", self.students[0].name),
                              ("MX Session", session.name), ("MX Session Execution", execution.name),
                              ("MX Teacher Hour Entry", hour), ("MX Teacher", self.teacher.name)):
            self.assertFalse(frappe.has_permission(doctype, "read", doc=frappe.get_doc(doctype, name)))
        self.assertTrue(pay.list_pay_sources(self.teacher.name, "2020-01-01", "2020-12-31"))
        self.assertTrue(pay.list_pay_teachers())

    def test_27_outsider_system_manager_and_docshare_do_not_bypass_payroll(self):
        hour = self.lesson()[2]
        sheet = self.approved([self.line(hour)])
        users = (self.user(), self.user("System Manager"), "Guest")
        for user in users:
            with self.subTest(user=user):
                frappe.set_user(user)
                with self.assertRaises(frappe.PermissionError):
                    pay.get_pay_balance(sheet)
                with self.assertRaises(frappe.PermissionError):
                    pay.get_source_summary(hour)
                self.assertFalse(frappe.has_permission("MX Teacher Pay Sheet", "read",
                                                       doc=frappe.get_doc("MX Teacher Pay Sheet", sheet)))
        frappe.set_user("Administrator")
        with self.assertRaises(frappe.PermissionError):
            frappe.get_doc({"doctype": "DocShare", "share_doctype": "MX Teacher Pay Sheet",
                            "share_name": sheet, "user": users[0], "read": 1}).insert()

    def test_28_request_key_content_and_money_input_validation(self):
        hour = self.lesson()[2]
        key = uuid.uuid4().hex
        first = self.sheet([self.line(hour)], request_id=key)
        self.assertEqual(self.sheet([self.line(hour)], request_id=key), first)
        with self.assertRaises(frappe.ValidationError):
            self.sheet([self.line(hour, 5)], request_id=key)
        for value in ("NaN", "Infinity", "1.001", "-1"):
            with self.subTest(value=value), self.assertRaises(frappe.ValidationError):
                sheet = self.sheet([self.fixed(work_key=self.batch + str(value), amount=value)])
                pay.check_sheet(sheet)

    def test_29_concurrent_same_creation_and_check_have_one_source_claim(self):
        hour = self.lesson()[2]
        args = {"teacher": self.teacher.name, "from_date": "2020-01-01", "to_date": "2020-12-31",
                "lines": [self.line(hour)], "request_id": uuid.uuid4().hex}
        results = self.race("pay_create", self.teacher.name, args)
        sheet = pay.create_sheet(**args)
        self.assertEqual(sheet, results[0]["name"])
        self.commit()
        self.race("pay_check", sheet, {"sheet": sheet})
        pay.check_sheet(sheet)
        self.assertEqual(frappe.db.count("MX Teacher Pay Sheet", {"demo_batch": self.batch}), 1)
        duplicate = self.sheet([self.line(hour)])
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(duplicate)

    def test_30_concurrent_approve_and_payment_are_idempotent(self):
        hour = self.lesson()[2]
        sheet = self.sheet([self.line(hour)])
        pay.check_sheet(sheet)
        self.race("pay_approve", sheet, {"sheet": sheet})
        self.assertEqual(pay.approve_sheet(sheet), sheet)
        payment_args = {"sheet": sheet, "paid_at": str(now_datetime()), "payment_method": "现金",
                        "reference_no": "并发付款同一凭据", "request_id": uuid.uuid4().hex}
        results = self.race("pay_payment", sheet, payment_args)
        self.assertEqual(pay.record_pay_payment(**payment_args), results[0]["name"])
        self.assertEqual(frappe.db.count("MX Teacher Pay Payment", {"sheet": sheet}), 1)
        self.assertEqual(self.money(sheet, "net_paid"), 100)

    def test_31_distinct_concurrent_sheets_cannot_both_claim_same_source(self):
        hour = self.lesson()[2]
        first, second = self.sheet([self.line(hour)]), self.sheet([self.line(hour)])
        results = self.race("pay_check", first, {"sheet": first}, "pay_check", second, {"sheet": second})
        self.assertFalse(results[1]["ok"], results)
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(second)
        pay.approve_sheet(first)
        self.assertEqual(self.money(first), 100)

    def test_32_approve_cancel_race_cannot_create_untracked_payable(self):
        _, execution, hour = self.lesson()
        sheet = self.sheet([self.line(hour)])
        pay.check_sheet(sheet)
        results = self.race("pay_approve", sheet, {"sheet": sheet},
                            "execution_cancel", execution.name, {})
        if not results[1]["ok"]:
            frappe.get_doc("MX Session Execution", execution.name).cancel()
        self.assertEqual(frappe.db.get_value("MX Session Execution", execution.name, "docstatus"), 2)
        self.assertEqual(len(self.adjustments(sheet)), 1)
        with self.assertRaises(frappe.ValidationError):
            self.payment(sheet)

    def test_33_payment_cancel_race_preserves_cash_and_adjustment_once(self):
        _, execution, hour = self.lesson()
        sheet = self.approved([self.line(hour)])
        args = {"sheet": sheet, "paid_at": str(now_datetime()), "payment_method": "现金",
                "reference_no": "付款先取得锁", "request_id": uuid.uuid4().hex}
        results = self.race("pay_payment", sheet, args, "execution_cancel", execution.name, {})
        if not results[1]["ok"]:
            frappe.get_doc("MX Session Execution", execution.name).cancel()
        self.assertEqual(self.money(sheet, "net_paid"), 100)
        self.assertEqual(len(self.adjustments(sheet)), 1)
        self.assertEqual(pay.get_pay_balance(sheet)["status"], "待处理调整")

    def test_34_cancel_wins_race_blocks_payment(self):
        _, execution, hour = self.lesson()
        sheet = self.approved([self.line(hour)])
        args = {"sheet": sheet, "paid_at": str(now_datetime()), "payment_method": "现金",
                "reference_no": "取消先取得锁", "request_id": uuid.uuid4().hex}
        results = self.race("execution_cancel", execution.name, {}, "pay_payment", sheet, args)
        self.assertFalse(results[1]["ok"], results)
        with self.assertRaises(frappe.ValidationError):
            pay.record_pay_payment(**args)
        self.assertEqual(self.money(sheet, "net_paid"), 0)
        self.assertEqual(frappe.db.count("MX Teacher Pay Payment", {"sheet": sheet}), 0)

    def test_35_adjustment_failure_rolls_back_source_cancellation(self):
        _, execution, hour = self.lesson()
        sheet = self.approved([self.line(hour)])
        frappe.db.savepoint("m4b_cancel_fault")
        original_insert = frappe.model.document.Document.insert

        def insert_or_fail(doc, *args, **kwargs):
            if doc.doctype == "MX Teacher Pay Adjustment":
                raise frappe.ValidationError("注入工资调整写入失败")
            return original_insert(doc, *args, **kwargs)

        with patch("frappe.model.document.Document.insert", insert_or_fail):
            with self.assertRaises(frappe.ValidationError):
                execution.cancel()
        frappe.db.rollback(save_point="m4b_cancel_fault")
        execution.reload()
        self.assertEqual(execution.docstatus, 1)
        self.assertEqual(frappe.db.count("MX Teacher Hour Entry", {"reversal_of": hour}), 0)
        self.assertEqual(self.adjustments(sheet), [])
        execution.cancel()
        self.assertEqual(len(self.adjustments(sheet)), 1)
        self.assertEqual(frappe.db.count("MX Teacher Hour Entry", {"reversal_of": hour}), 1)

    def test_36_actual_finance_role_can_approve_and_register_without_source_access(self):
        hour = self.lesson()[2]
        sheet = self.sheet([self.line(hour)])
        scheduler, finance = self.user("Meixin Scheduler"), self.user("Meixin Finance")
        frappe.set_user(scheduler)
        pay.check_sheet(sheet)
        frappe.set_user(finance)
        self.assertEqual(pay.approve_sheet(sheet), sheet)
        payment = self.payment(sheet)
        self.assertEqual(Decimal(str(frappe.db.get_value("MX Teacher Pay Payment", payment, "amount"))), 100)
        self.assertEqual(pay.get_pay_balance(sheet)["status"], "已付款")
        self.assertFalse(frappe.has_permission("MX Teacher Hour Entry", "read"))

    def test_37_cancel_first_race_blocks_approval(self):
        _, execution, hour = self.lesson()
        sheet = self.sheet([self.line(hour)])
        pay.check_sheet(sheet)
        results = self.race("execution_cancel", execution.name, {}, "pay_approve", sheet, {"sheet": sheet})
        self.assertFalse(results[1]["ok"], results)
        with self.assertRaises(frappe.ValidationError):
            pay.approve_sheet(sheet)
        self.assertEqual(frappe.db.get_value("MX Teacher Pay Sheet", sheet, "docstatus"), 0)

    def test_38_concurrent_whole_day_and_scoped_claims_do_not_overlap(self):
        first = self.lesson()[2]
        second = self.lesson()[2]
        whole = self.sheet([self.fixed(hours=[first, second], scope="全天")])
        scoped = self.sheet([self.fixed(hours=[second], scope="指定")])
        results = self.race("pay_check", whole, {"sheet": whole}, "pay_check", scoped, {"sheet": scoped})
        self.assertFalse(results[1]["ok"], results)
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(scoped)

    def test_39_concurrent_payment_reversal_creates_one_controlled_inverse(self):
        sheet = self.approved([self.line(self.lesson()[2])])
        payment = self.payment(sheet)
        args = {"payment": payment, "reason": "虚构付款误登记，经财务核实",
                "request_id": uuid.uuid4().hex}
        results = self.race("pay_reverse", payment, args)
        self.assertEqual(pay.reverse_pay_payment(**args), results[0]["name"])
        self.assertEqual(frappe.db.count("MX Teacher Pay Payment", {"sheet": sheet}), 2)
        self.assertEqual(self.money(sheet, "net_paid"), 0)
        self.payment(sheet)
        self.assertEqual(self.money(sheet, "net_paid"), 100)

    def test_40_report_and_csv_excel_export_follow_payroll_read_permissions(self):
        from frappe.desk import query_report
        from meixin_admin.meixin_admin.report.教师工资结算概览.教师工资结算概览 import execute

        sheet = self.approved([self.line(self.lesson()[2])])
        allowed = [self.user(role) for role in ("Meixin Scheduler", "Meixin Manager", "Meixin Finance")]
        denied = [self.user(), self.user("System Manager"), "Guest"]
        original_form, original_response = frappe.local.form_dict, frappe.local.response
        try:
            for user in allowed:
                frappe.set_user(user)
                filters = {"teacher": self.teacher.name, "from_date": "2020-01-01", "to_date": "2020-12-31"}
                self.assertIn(sheet, [row["sheet"] for row in execute(filters)[1]])
                for file_format in ("CSV", "Excel"):
                    with self.subTest(user=user, format=file_format):
                        frappe.local.form_dict = frappe._dict({
                            "report_name": "教师工资结算概览", "file_format_type": file_format,
                            "filters": json.dumps(filters),
                        })
                        frappe.local.response = frappe._dict()
                        query_report.export_query()
                        content = frappe.local.response.get("filecontent")
                        self.assertTrue(content)
                        if file_format == "CSV":
                            self.assertIn(sheet, content.decode("utf-8-sig"))
                        else:
                            self.assertEqual(content[:2], b"PK")
            for user in denied:
                frappe.set_user(user)
                with self.assertRaises(frappe.PermissionError):
                    execute()
                frappe.local.form_dict = frappe._dict({
                    "report_name": "教师工资结算概览", "file_format_type": "CSV", "filters": "{}",
                })
                frappe.local.response = frappe._dict()
                with self.assertRaises(frappe.PermissionError):
                    query_report.export_query()
        finally:
            frappe.local.form_dict, frappe.local.response = original_form, original_response

    def test_41_rule_ranges_must_be_complete_nonoverlapping_and_amounts_precise(self):
        cases = (
            [{"min_students": 1, "max_students": 4, "amount": 100},
             {"min_students": 4, "max_students": 0, "amount": 120}],
            [{"min_students": 1, "max_students": 3, "amount": 100},
             {"min_students": 5, "max_students": 0, "amount": 120}],
            [{"min_students": 1, "max_students": 0, "amount": "100.001"}],
        )
        for rates in cases:
            with self.subTest(rates=rates), self.assertRaises(frappe.ValidationError):
                frappe.get_doc({"doctype": "MX Teacher Pay Rule", "effective_from": "2020-02-01",
                                "currency": "CNY", "rates": rates, "demo_batch": self.batch}).insert().submit()

    def test_42_request_validation_rejects_invalid_types_and_missing_bounds(self):
        for lines in ("not json", "{}", {"kind": "课程"}, None):
            with self.subTest(lines=lines), self.assertRaises(frappe.ValidationError):
                self.sheet(lines)
        hour = self.lesson()[2]
        for size in ([], {}, True, "4.5", -1):
            with self.subTest(size=size), self.assertRaises(frappe.ValidationError):
                sheet = self.sheet([self.line(hour, size)])
                pay.check_sheet(sheet)
        for malformed in (self.line(hour, kind=[]), self.line(hour, kind={}),
                          self.fixed(scope=[]), self.fixed(scope={})):
            with self.subTest(line=malformed), self.assertRaises(frappe.ValidationError):
                self.sheet([malformed])
        with self.assertRaises(frappe.ValidationError):
            pay.create_sheet(self.teacher.name, "2020-02-01", "2020-01-01",
                             [self.line(hour)], uuid.uuid4().hex)

    def test_43_session_revision_chain_keeps_original_payroll_identity(self):
        session, execution, hour = self.lesson()
        sheet = self.approved([self.line(hour)])
        execution.cancel()
        session.cancel()
        revised_session = frappe.copy_doc(session)
        revised_session.amended_from, revised_session.docstatus = session.name, 0
        revised_session.insert().submit()
        revised_execution = frappe.get_doc({
            "doctype": "MX Session Execution", "session": revised_session.name,
            "attendance": [{"student": row.student, "attendance_status": "到课"}
                           for row in revised_session.students],
        }).insert().submit()
        revised_hour = confirm_teaching(revised_execution.name, revised_session.start_at, revised_session.end_at)
        duplicate = self.sheet([self.line(revised_hour)])
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(duplicate)
        self.assertEqual(self.money(sheet, "total_amount"), 100)

    def test_44_cancelled_planned_course_creates_no_payable_source(self):
        session = frappe.get_doc({
            "doctype": "MX Session", "teacher": self.teacher.name, "course": self.course.name,
            "room": self.room.name, "start_at": self.day, "end_at": self.day + timedelta(minutes=90),
            "students": [{"student": student.name} for student in self.students[:4]],
            "demo_batch": self.batch,
        }).insert().submit()
        session.cancel()
        self.assertEqual(pay.list_pay_sources(self.teacher.name, "2020-01-01", "2020-12-31"), [])
        sheet = self.sheet([self.fixed(scope="全天")])
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(sheet)

    def test_45_readonly_finance_endpoints_never_take_row_locks(self):
        hour = self.lesson()[2]
        sheet = self.approved([self.line(hour)])
        finance = self.user("Meixin Finance")
        frappe.set_user(finance)
        with patch.object(frappe.db, "sql", wraps=frappe.db.sql) as queries:
            self.assertTrue(pay.get_source_summary(hour))
            self.assertTrue(pay.list_pay_sources(self.teacher.name, "2020-01-01", "2020-12-31"))
            self.assertTrue(pay.list_pay_teachers())
            self.assertEqual(pay.get_pay_balance(sheet)["net_payable"], 100)
            self.assertTrue(pay.get_pay_report({"teacher": self.teacher.name})[1])
        self.assertTrue(queries.call_args_list)
        for call in queries.call_args_list:
            self.assertNotIn("FOR UPDATE", str(call.args[0]).upper())

    def test_46_finance_can_read_cancelled_source_summary_without_private_data(self):
        session, execution, hour = self.lesson()
        finance = self.user("Meixin Finance")
        execution.cancel()
        frappe.set_user(finance)
        with patch.object(frappe.db, "sql", wraps=frappe.db.sql) as queries:
            summary = pay.get_source_summary(hour)
        self.assertFalse(summary["is_valid"])
        self.assertTrue(summary["reversal"])
        self.assertEqual(summary["source_root"], session.name)
        payload = json.dumps(summary, ensure_ascii=False, default=str)
        for secret in ("guardian_phone", "attendance", "headcount_basis", "exception_reason",
                       "00000000000", *(student.student_name for student in self.students),
                       *(student.name for student in self.students)):
            self.assertNotIn(secret, payload)
        for call in queries.call_args_list:
            self.assertNotIn("FOR UPDATE", str(call.args[0]).upper())
        for doctype, name in (("MX Session", session.name), ("MX Session Execution", execution.name),
                              ("MX Teacher Hour Entry", hour)):
            self.assertFalse(frappe.has_permission(doctype, "read", doc=frappe.get_doc(doctype, name)))

    def test_47_same_work_identity_cannot_be_rebilled_on_another_date(self):
        work_key = "Lesson-Plan-" + self.batch
        first = self.approved([self.fixed(work_key=work_key, amount=66)])
        duplicate = self.sheet([self.fixed(work_key=work_key.upper(), date="2020-01-16", amount=66)])
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(duplicate)
        separate = self.approved([self.fixed(work_key=work_key + "-separate", amount=77)])
        self.assertEqual(self.money(first), 66)
        self.assertEqual(self.money(separate), 77)

    def test_48_submitted_rule_is_immutable_and_external_edit_cannot_rewrite_old_bill(self):
        sheet = self.approved([self.line(self.lesson()[2])])
        before = frappe.get_doc("MX Teacher Pay Sheet", sheet).as_json()
        rule = frappe.get_doc("MX Teacher Pay Rule", self.rule.name)
        rule.rates[0].amount = 150
        with self.assertRaises(frappe.ValidationError):
            rule.save()
        # Even an external SQL edit cannot retroactively replace an approved snapshot.
        frappe.db.set_value("MX Teacher Pay Rate", self.rule.rates[0].name, "amount", 150)
        self.assertEqual(self.money(sheet), 100)
        self.assertEqual(frappe.get_doc("MX Teacher Pay Sheet", sheet).as_json(), before)

    def test_49_business_roles_cannot_write_child_rows_outside_parent_operations(self):
        sheet = self.approved([self.line(self.lesson()[2])])
        parent_before = frappe.get_doc("MX Teacher Pay Sheet", sheet).as_json()
        children = (
            ("MX Teacher Pay Rate", self.rule.rates[0].name),
            ("MX Teacher Pay Line", frappe.get_doc("MX Teacher Pay Sheet", sheet).lines[0].name),
            ("MX Teacher Pay Coverage", frappe.db.get_value("MX Teacher Pay Coverage", {"parent": sheet})),
        )
        users = [self.user(role) for role in ("Meixin Scheduler", "Meixin Finance")]
        for user in users:
            frappe.set_user(user)
            for doctype, name in children:
                with self.subTest(user=user, doctype=doctype):
                    child = frappe.get_doc(doctype, name)
                    with self.assertRaises(frappe.PermissionError):
                        frappe.copy_doc(child).insert()
                    child.idx += 1
                    with self.assertRaises(frappe.PermissionError):
                        child.save()
                    with self.assertRaises(frappe.PermissionError):
                        frappe.delete_doc(doctype, name)
        self.assertEqual(frappe.get_doc("MX Teacher Pay Sheet", sheet).as_json(), parent_before)
        self.assertEqual(self.money(sheet), 100)

    def test_50_large_report_and_source_lists_fail_explicitly_instead_of_truncating(self):
        finance = self.user("Meixin Finance")
        frappe.set_user(finance)
        frappe.get_roles(finance)
        filters = {"teacher": self.teacher.name, "from_date": "2020-01-01", "to_date": "2020-12-31"}
        with patch.object(frappe, "get_list", return_value=["sentinel"] * 1001) as query:
            with self.assertRaisesRegex(frappe.ValidationError, "1000.*不会截断"):
                pay.get_pay_report(filters)
        query.assert_called_once_with(
            "MX Teacher Pay Sheet", pluck="name", order_by="creation desc", limit_page_length=1001,
            filters={"teacher": self.teacher.name, "to_date": [">=", datetime(2020, 1, 1).date()],
                     "from_date": ["<=", datetime(2020, 12, 31).date()]},
        )
        with patch.object(frappe.db, "sql", return_value=[("source",)] * 501) as query:
            with patch.object(pay, "_source") as read_source:
                with self.assertRaisesRegex(frappe.ValidationError, "500.*不会截断"):
                    pay.list_pay_sources(self.teacher.name, "2020-01-01", "2020-12-31")
                read_source.assert_not_called()
        query.assert_called_once()
        self.assertIn("LIMIT 501", query.call_args.args[0])
        self.assertEqual(query.call_args.args[1],
                         (self.teacher.name, datetime(2020, 1, 1).date(), datetime(2021, 1, 1).date()))

    def test_51_adjustment_cannot_overflow_existing_net_and_remains_unapproved(self):
        from meixin_admin.money import MAX_CNY

        _, execution, hour = self.lesson()
        sheet = self.approved([self.line(hour)])
        execution.cancel()
        adjustment = self.adjustments(sheet)[0].name
        finance = self.user("Meixin Finance")
        frappe.set_user(finance)
        original = frappe.get_doc("MX Teacher Pay Adjustment", adjustment).as_json()
        with self.assertRaisesRegex(frappe.ValidationError, "范围"):
            pay.approve_adjustment(adjustment, str(MAX_CNY), "单笔合法但净应付超出上限")
        unchanged = frappe.get_doc("MX Teacher Pay Adjustment", adjustment)
        self.assertEqual(unchanged.as_json(), original)
        self.assertEqual(unchanged.docstatus, 0)
        self.assertEqual(Decimal(str(unchanged.amount)), 0)
        self.assertEqual(self.money(sheet), 100)

    def test_52_payroll_currency_metadata_documents_report_and_excel_are_cny(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from frappe.desk import query_report
        from frappe.utils.xlsxutils import XLSXStyleBuilder
        from meixin_admin.payroll_permissions import PAYROLL_DOCTYPES

        for doctype in PAYROLL_DOCTYPES:
            meta = frappe.get_meta(doctype)
            for field in meta.fields:
                if field.fieldtype == "Currency":
                    with self.subTest(doctype=doctype, field=field.fieldname):
                        self.assertEqual(field.options, "currency")
                        self.assertEqual(meta.get_field("currency").default, "CNY")
        _, execution, hour = self.lesson()
        sheet = self.approved([self.line(hour)])
        payment = self.payment(sheet)
        execution.cancel()
        adjustment = self.adjustments(sheet)[0].name
        for doctype, name in (("MX Teacher Pay Sheet", sheet), ("MX Teacher Pay Rule", self.rule.name),
                              ("MX Teacher Pay Payment", payment), ("MX Teacher Pay Adjustment", adjustment)):
            self.assertEqual(frappe.get_doc(doctype, name).currency, "CNY")
        self.assertTrue(all(row.currency == "CNY" for row in frappe.get_doc("MX Teacher Pay Rule", self.rule.name).rates))
        self.assertTrue(all(row.currency == "CNY" for row in frappe.get_doc("MX Teacher Pay Sheet", sheet).lines))
        columns, rows = pay.get_pay_report({"teacher": self.teacher.name})
        self.assertTrue(rows)
        self.assertTrue(all(row["currency"] == "CNY" for row in rows))
        self.assertTrue(all(col.get("options") == "currency" for col in columns if col["fieldtype"] == "Currency"))
        original_form, original_response = frappe.local.form_dict, frappe.local.response
        try:
            frappe.local.form_dict = frappe._dict({"report_name": "教师工资结算概览", "file_format_type": "Excel",
                                                  "filters": json.dumps({"teacher": self.teacher.name})})
            frappe.local.response = frappe._dict()
            query_report.export_query()
            workbook = load_workbook(BytesIO(frappe.local.response.filecontent))
            exported = next(row for row in workbook.active.iter_rows() if row[0].value == sheet)
            amount_column = next(i for i, col in enumerate(columns) if col["fieldname"] == "amount")
            self.assertEqual(exported[amount_column].value, 100)
            self.assertEqual(exported[amount_column].number_format,
                             XLSXStyleBuilder.get_number_format("Currency", "CNY"))
        finally:
            frappe.local.form_dict, frappe.local.response = original_form, original_response

    def test_53_scoped_fixed_fee_can_cover_two_days_without_bypassing_day_claims(self):
        first = self.lesson(start=self.day)[2]
        second = self.lesson(start=self.day + timedelta(days=1))[2]
        sheet = self.approved([self.fixed(hours=[first, second], amount=300)])
        approved = frappe.get_doc("MX Teacher Pay Sheet", sheet)
        self.assertEqual(len(approved.lines), 1)
        self.assertEqual(self.money(sheet), 300)
        covered = {row.hour_entry: str(row.pay_date) for row in approved.coverage if row.hour_entry}
        self.assertEqual(covered, {first: "2020-01-15", second: "2020-01-16"})
        for line in (self.line(second), self.fixed(scope="全天", date="2020-01-16")):
            with self.subTest(line=line), self.assertRaises(frappe.ValidationError):
                pay.check_sheet(self.sheet([line]))
        with self.assertRaisesRegex(frappe.ValidationError, "结算范围"):
            pay.create_sheet(self.teacher.name, "2020-01-15", "2020-01-15",
                             [self.fixed(hours=[first, second])], uuid.uuid4().hex)
        self.teacher = self.master("MX Teacher", teacher_name="跨日反向覆盖虚构教师")
        earlier = self.lesson(start=self.day + timedelta(days=2))[2]
        later = self.lesson(start=self.day + timedelta(days=3))[2]
        self.approved([self.fixed(scope="全天", date="2020-01-18", amount=300)])
        duplicate = self.sheet([self.fixed(hours=[earlier, later], date="2020-01-17", amount=300)])
        with self.assertRaises(frappe.ValidationError):
            pay.check_sheet(duplicate)

    def test_54_missing_special_fixed_and_adjustment_amounts_do_not_become_implicit_zero(self):
        _, execution, hour = self.lesson()
        for kind in ("特殊", "固定"):
            for amount in (None, "", "missing"):
                line = (self.line(hour, kind="特殊", exception_reason="特殊金额待明确", amount=amount)
                        if kind == "特殊" else self.fixed(work_key="未明确金额-" + self.batch, amount=amount))
                if amount == "missing":
                    del line["amount"]
                with self.subTest(kind=kind, amount=amount):
                    sheet = self.sheet([line])
                    self.assertEqual(pay.get_pay_balance(sheet)["status"], "待核对")
                    with self.assertRaisesRegex(frappe.ValidationError, "明确填写金额"):
                        pay.check_sheet(sheet)
        zero_special = self.approved([self.line(hour, kind="特殊", amount=0, exception_reason="明确本课零应付")])
        zero_fixed = self.approved([self.fixed(work_key="明确零费用-" + self.batch, amount=0)])
        for sheet in (zero_special, zero_fixed):
            self.assertEqual(pay.get_pay_balance(sheet)["status"], "无需付款")
            with self.assertRaises(frappe.ValidationError):
                self.payment(sheet)
            self.assertEqual(frappe.db.count("MX Teacher Pay Payment", {"sheet": sheet}), 0)
        execution.cancel()
        adjustment = self.adjustments(zero_special)[0].name
        before = frappe.get_doc("MX Teacher Pay Adjustment", adjustment).as_json()
        for amount in (None, ""):
            with self.subTest(adjustment_amount=amount), self.assertRaises(frappe.ValidationError):
                pay.approve_adjustment(adjustment, amount, "未明确金额不能视为零决定")
            self.assertEqual(frappe.get_doc("MX Teacher Pay Adjustment", adjustment).as_json(), before)
        pay.approve_adjustment(adjustment, 0, "财务明确零调整")
        self.assertEqual(pay.get_pay_balance(zero_special)["status"], "无需付款")
